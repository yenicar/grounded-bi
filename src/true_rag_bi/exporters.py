from __future__ import annotations

import io
from html import escape

from .utils import utc_now_iso


def export_answer_pdf(payload: dict) -> bytes:
    try:
        from weasyprint import HTML
    except ImportError as exc:
        raise RuntimeError("PDF export requires weasyprint. Install the export dependencies first.") from exc

    citations = "".join(
        f"<li><strong>{escape(item['doc_name'])}</strong> ({escape(item['page_range'])})<br>{escape(item['excerpt'])}</li>"
        for item in payload.get("citations", [])
    )
    html = f"""
    <html>
      <body style="font-family: Arial, sans-serif; padding: 32px;">
        <h1>{escape(payload.get('title') or 'GroundedBI Export')}</h1>
        <p><strong>Question:</strong> {escape(payload.get('question') or '')}</p>
        <h2>Answer</h2>
        <div>{escape(payload.get('answer') or '').replace(chr(10), '<br>')}</div>
        <h2>Citations</h2>
        <ol>{citations}</ol>
        <p style="margin-top: 24px; color: #666;">Generated {escape(payload.get('generated_at') or utc_now_iso())}</p>
      </body>
    </html>
    """
    return HTML(string=html).write_pdf()


def export_answer_pptx(payload: dict) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("PPTX export requires python-pptx. Install the export dependencies first.") from exc

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.0), Inches(0.8))
    title_frame = title_box.text_frame
    title_frame.text = payload.get("title") or "GroundedBI Export"
    title_frame.paragraphs[0].font.size = Pt(24)
    title_frame.paragraphs[0].font.bold = True

    question_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(12.0), Inches(0.8))
    question_box.text_frame.text = f"Question: {payload.get('question') or ''}"

    answer_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(7.4), Inches(4.5))
    answer_frame = answer_box.text_frame
    answer_frame.word_wrap = True
    answer_frame.text = payload.get("answer") or ""

    citation_box = slide.shapes.add_textbox(Inches(8.1), Inches(2.0), Inches(4.3), Inches(4.5))
    citation_frame = citation_box.text_frame
    citation_frame.text = "Citations"
    citation_frame.paragraphs[0].font.bold = True
    for citation in payload.get("citations", [])[:5]:
        paragraph = citation_frame.add_paragraph()
        paragraph.text = f"{citation['doc_name']} ({citation['page_range']})"
        paragraph.level = 0
        paragraph.font.size = Pt(12)

    footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(12.0), Inches(0.4))
    footer_box.text_frame.text = f"Generated {payload.get('generated_at') or utc_now_iso()}"

    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
