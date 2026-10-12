from __future__ import annotations

from pathlib import Path
from typing import Any

from .cache import CacheStore
from .config import AppConfig
from .corpus import infer_policy_category, infer_source_level, list_source_files
from .metadata import extract_document_metadata
from .utils import ensure_directory, logger, normalize_text, sha256_file, sha256_text, utc_now_iso, write_json

_INGEST_PIPELINE_VERSION = "v6_multiformat_docx_xlsx_csv"
_SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv"}


def _extract_document(file_path: Path) -> tuple[list[str], list[str]]:
    """Route a document to the correct extractor based on file extension."""
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_pages(file_path)
    if suffix == ".docx":
        return _extract_docx_pages(file_path)
    if suffix in {".xlsx", ".csv"}:
        return _extract_spreadsheet_pages(file_path)
    raise RuntimeError(f"Unsupported file type: {suffix}. Supported: {sorted(_SUPPORTED_EXTENSIONS)}")


def _extract_pdf_pages(pdf_path: Path) -> tuple[list[str], list[str]]:
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise RuntimeError("PyPDF2 is required for ingestion. Install project dependencies first.") from exc

    reader = PdfReader(str(pdf_path))
    pages: list[str] = []
    warnings: list[str] = []

    # Open pdfplumber ONCE for the entire document to avoid O(n) file opens
    plumber_pdf = _open_pdfplumber(pdf_path)

    for page_number, page in enumerate(reader.pages, start=1):
        raw_text = page.extract_text() or ""
        if not raw_text:
            raw_text = _extract_page_pdfminer(pdf_path, page_number)

        table_blocks = _extract_page_tables_from_open_pdf(plumber_pdf, page_number)
        page_text = _merge_page_text_and_tables(raw_text, table_blocks)
        normalized_page = normalize_text(page_text)

        if table_blocks:
            warnings.append(f"page_{page_number}_tables_extracted_{len(table_blocks)}")
        if not normalized_page:
            warnings.append(f"page_{page_number}_no_extractable_text")

        pages.append(normalized_page)

    if plumber_pdf is not None:
        try:
            plumber_pdf.close()
        except Exception:
            pass

    return pages, warnings


def _open_pdfplumber(pdf_path: Path):
    """Open a pdfplumber PDF handle. Returns None if pdfplumber is unavailable."""
    try:
        import pdfplumber
        return pdfplumber.open(str(pdf_path))
    except ImportError:
        return None
    except Exception:
        return None


def _extract_page_pdfminer(pdf_path: Path, page_number: int) -> str:
    try:
        import io

        from pdfminer.high_level import extract_text_to_fp
        from pdfminer.layout import LAParams
    except ImportError:
        return ""
    except Exception:
        return ""

    try:
        output = io.StringIO()
        with pdf_path.open("rb") as handle:
            extract_text_to_fp(
                handle,
                output,
                laparams=LAParams(),
                page_numbers={page_number - 1},
                output_type="text",
                codec="utf-8",
            )
        return output.getvalue()
    except Exception:
        return ""


def _extract_page_tables_from_open_pdf(plumber_pdf, page_number: int) -> list[str]:
    """Extract tables from an already-open pdfplumber PDF object."""
    if plumber_pdf is None:
        return []

    strategies = [
        {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
        {"vertical_strategy": "text", "horizontal_strategy": "text"},
    ]
    table_markdowns: list[str] = []
    seen_hashes: set[str] = set()

    try:
        if page_number - 1 >= len(plumber_pdf.pages):
            return []
        page = plumber_pdf.pages[page_number - 1]
        for settings in strategies:
            strategy_tables: list[str] = []
            for table in page.extract_tables(table_settings=settings) or []:
                markdown_table = _format_markdown_table(table)
                if not markdown_table:
                    continue
                table_hash = sha256_text(markdown_table)
                if table_hash in seen_hashes:
                    continue
                seen_hashes.add(table_hash)
                strategy_tables.append(markdown_table)
            if strategy_tables:
                table_markdowns.extend(strategy_tables)
                break
    except Exception:
        return []

    return table_markdowns


def _normalize_table_cell(value: Any) -> str:
    if value is None:
        return ""
    text = normalize_text(str(value))
    text = text.replace("|", "\\|")
    return text


def _looks_numeric(value: str) -> bool:
    if not value:
        return False
    digits = sum(character.isdigit() for character in value)
    return digits >= 2


def _looks_like_structured_table(rows: list[list[Any]] | None) -> bool:
    if not rows:
        return False
    normalized_rows = [
        [_normalize_table_cell(cell) for cell in row]
        for row in rows
        if row and any(_normalize_table_cell(cell) for cell in row)
    ]
    if len(normalized_rows) < 2:
        return False

    width = max(len(row) for row in normalized_rows)
    if width < 2:
        return False

    header = normalized_rows[0]
    header_filled = sum(1 for cell in header if cell)
    if header_filled < 2:
        return False

    if sum(_looks_numeric(cell) for cell in header if cell) > 1:
        return False

    body = normalized_rows[1:]
    dense_body_rows = sum(1 for row in body if sum(1 for cell in row if cell) >= 2)
    numeric_cells = sum(_looks_numeric(cell) for row in body for cell in row if cell)
    numeric_rows = sum(1 for row in body if sum(_looks_numeric(cell) for cell in row if cell) >= 1)
    sparse_rows = sum(1 for row in body if sum(1 for cell in row if cell) <= 1)
    minimum_numeric_cells = 2 if width == 2 else 3
    return (
        dense_body_rows >= 2
        and numeric_cells >= minimum_numeric_cells
        and numeric_rows >= 2
        and sparse_rows <= 1
    )


def _format_markdown_table(rows: list[list[Any]] | None) -> str:
    if not _looks_like_structured_table(rows):
        return ""

    normalized_rows = [
        [_normalize_table_cell(cell) for cell in row]
        for row in rows
        if row and any(_normalize_table_cell(cell) for cell in row)
    ]
    if len(normalized_rows) < 2:
        return ""

    width = max(len(row) for row in normalized_rows)
    padded_rows = [row + [""] * (width - len(row)) for row in normalized_rows]
    header = padded_rows[0]
    separator = ["---"] * width
    body = padded_rows[1:]

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def _merge_page_text_and_tables(raw_text: str, table_blocks: list[str]) -> str:
    normalized_text = normalize_text(raw_text)
    if not table_blocks:
        return normalized_text

    sections: list[str] = []
    if normalized_text:
        sections.append(normalized_text)
    for index, table in enumerate(table_blocks, start=1):
        sections.append(f"### Extracted Table {index}\n{table}")
    return "\n\n".join(section for section in sections if section.strip())


def _extract_docx_pages(docx_path: Path) -> tuple[list[str], list[str]]:
    """Extract text and tables from a .docx file. Each logical section becomes a page."""
    try:
        import docx  # python-docx
    except ImportError:
        return (["[[DOCX_EXTRACTION_REQUIRES_PYTHON_DOCX]]"], ["python_docx_not_installed"])

    warnings: list[str] = []
    pages: list[str] = []
    current_page_lines: list[str] = []
    table_count = 0

    try:
        document = docx.Document(str(docx_path))
    except Exception as exc:
        return ([f"[[DOCX_EXTRACTION_FAILED: {exc}]]"], ["docx_extraction_failed"])

    # Iterate child elements in document order (paragraphs and tables interleaved)
    for block in document.element.body:
        tag = block.tag.split("}")[-1] if "}" in block.tag else block.tag

        if tag == "p":
            import docx.oxml.ns as ns
            para = docx.text.paragraph.Paragraph(block, document)
            text = normalize_text(para.text)
            if not text:
                continue
            # Heading styles act as page breaks for chunking clarity
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                if current_page_lines:
                    pages.append("\n".join(current_page_lines))
                    current_page_lines = []
            current_page_lines.append(text)

        elif tag == "tbl":
            import docx.table as tbl_mod
            table = docx.table.Table(block, document)
            rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
            markdown = _format_markdown_table(rows)
            if markdown:
                table_count += 1
                if current_page_lines:
                    pages.append("\n".join(current_page_lines))
                    current_page_lines = []
                pages.append(f"### Extracted Table {table_count}\n{markdown}")
                warnings.append(f"table_{table_count}_extracted_from_docx")

    if current_page_lines:
        pages.append("\n".join(current_page_lines))

    if not pages:
        warnings.append("docx_no_extractable_content")
        pages = ["[[NO_EXTRACTED_TEXT]]"]

    return pages, warnings


def _extract_spreadsheet_pages(file_path: Path) -> tuple[list[str], list[str]]:
    """Extract XLSX or CSV to Markdown tables. Each sheet/file becomes a page."""
    warnings: list[str] = []
    suffix = file_path.suffix.lower()

    if suffix == ".csv":
        return _extract_csv_pages(file_path, warnings)
    return _extract_xlsx_pages(file_path, warnings)


def _extract_csv_pages(file_path: Path, warnings: list[str]) -> tuple[list[str], list[str]]:
    try:
        import csv
        rows: list[list[str]] = []
        with file_path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append([str(cell) for cell in row])
        if not rows:
            warnings.append("csv_empty")
            return ["[[NO_EXTRACTED_TEXT]]"], warnings
        markdown = _format_markdown_table(rows)
        if not markdown:
            # Table failed heuristic — emit as plain text
            markdown = "\n".join(", ".join(row) for row in rows[:50])
            warnings.append("csv_plain_text_fallback")
        return [f"### {file_path.stem}\n{markdown}"], warnings
    except Exception as exc:
        warnings.append(f"csv_extraction_failed: {exc}")
        return ["[[NO_EXTRACTED_TEXT]]"], warnings


def _extract_xlsx_pages(file_path: Path, warnings: list[str]) -> tuple[list[str], list[str]]:
    try:
        import openpyxl
    except ImportError:
        return (["[[XLSX_EXTRACTION_REQUIRES_OPENPYXL]]"], ["openpyxl_not_installed"])

    try:
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    except Exception as exc:
        warnings.append(f"xlsx_load_failed: {exc}")
        return ["[[NO_EXTRACTED_TEXT]]"], warnings

    pages: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(["" if cell is None else str(cell) for cell in row])
        if not rows:
            warnings.append(f"xlsx_sheet_{sheet_name}_empty")
            continue
        markdown = _format_markdown_table(rows)
        if not markdown:
            markdown = "\n".join(", ".join(row) for row in rows[:50])
            warnings.append(f"xlsx_sheet_{sheet_name}_plain_text_fallback")
        pages.append(f"### Sheet: {sheet_name}\n{markdown}")
        warnings.append(f"xlsx_sheet_{sheet_name}_extracted_{len(rows)}_rows")

    if not pages:
        warnings.append("xlsx_no_extractable_sheets")
        pages = ["[[NO_EXTRACTED_TEXT]]"]

    return pages, warnings


def _build_markdown(doc_name: str, pages: list[str]) -> str:
    lines = [f"# {doc_name}", ""]
    for page_number, page_text in enumerate(pages, start=1):
        lines.append(f"## Page {page_number}")
        lines.append("")
        lines.append(page_text if page_text else "[[NO_EXTRACTED_TEXT]]")
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def ingest_documents(config: AppConfig, verbose: bool = False) -> dict[str, Any]:
    paths = config.paths
    ensure_directory(paths.normalized_md_dir)

    docs_dir = paths.raw_pdfs_dir
    ensure_directory(docs_dir)

    cache = CacheStore(paths.cache_dir)
    manifest = cache.load("ingestion_manifest", {"documents": [], "updated_at": None})
    previous_records = {
        record.get("source_relative_path") or record["doc_name"]: record
        for record in manifest.get("documents", [])
    }

    records: list[dict[str, Any]] = []
    cache_hits = 0
    cache_misses = 0

    source_files = list_source_files(docs_dir)

    for doc_id, source in enumerate(source_files, start=1):
        file_path = source["path"]
        source_relative_path = source["source_relative_path"]
        artifact_id = source["artifact_id"]
        file_hash = source["source_hash"]
        doc_name = source_relative_path
        markdown_name = f"{artifact_id}.md"
        markdown_path = paths.normalized_md_dir / markdown_name
        existing = previous_records.get(source_relative_path)

        if (
            existing
            and existing.get("source_pdf_hash") == file_hash
            and existing.get("ingest_pipeline_version") == _INGEST_PIPELINE_VERSION
            and markdown_path.exists()
        ):
            cache_hits += 1
            record = dict(existing)
            record["doc_id"] = doc_id
            record["last_ingested_at"] = utc_now_iso()
            records.append(record)
            if verbose:
                logger.info("[ingest] cache hit for %s", source_relative_path)
            continue

        cache_misses += 1
        pages, warnings = _extract_document(file_path)
        metadata = extract_document_metadata(source_relative_path, pages)
        markdown_text = _build_markdown(source_relative_path, pages)
        ensure_directory(markdown_path.parent)
        markdown_path.write_text(markdown_text, encoding="utf-8")
        text_sample = "\n".join(pages[:3])

        record = {
            "doc_id": doc_id,
            "doc_name": doc_name,
            "doc_format": file_path.suffix.lower().lstrip("."),
            "source_pdf": str(file_path.relative_to(paths.root)),
            "source_relative_path": source_relative_path,
            "artifact_id": artifact_id,
            "source_pdf_hash": file_hash,
            "normalized_md_path": str(markdown_path.relative_to(paths.root)),
            "page_count": len(pages),
            "content_hash": sha256_text(markdown_text),
            "warnings": warnings,
            "doc_date": metadata.get("doc_date"),
            "doc_author": metadata.get("doc_author"),
            "doc_type": metadata.get("doc_type"),
            "source_level": infer_source_level(source_relative_path),
            "policy_category": infer_policy_category(source_relative_path, text_sample),
            "effective_date": metadata.get("doc_date"),
            "ingest_pipeline_version": _INGEST_PIPELINE_VERSION,
            "last_ingested_at": utc_now_iso(),
        }
        records.append(record)
        if verbose:
            logger.info("[ingest] extracted %s (%s pages)", source_relative_path, len(pages))

    payload = {
        "documents": records,
        "updated_at": utc_now_iso(),
        "cache_stats": {"hits": cache_hits, "misses": cache_misses},
    }
    cache.save("ingestion_manifest", payload)
    write_json(paths.corpus_dir / "ingestion_summary.json", payload)
    return payload
